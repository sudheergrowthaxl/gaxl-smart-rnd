"""Export derived rules to Excel or text."""

from pathlib import Path


def _parse_rule_line(line: str) -> tuple[str, str, str, str, str] | None:
    """Parse a tab-separated rule line into (Entity, Attribute, Normalization, Rule description, Few Shot examples)."""
    line = line.strip()
    if not line:
        return None
    parts = line.split("\t")
    if len(parts) >= 5:
        return (
            parts[0].strip(),
            parts[1].strip(),
            parts[2].strip(),
            parts[3].strip(),
            parts[4].strip(),
        )
    if len(parts) >= 4:
        return (parts[0].strip(), parts[1].strip(), parts[2].strip(), parts[3].strip(), "")
    if len(parts) == 1 and (line.startswith("Contactors") or "Normalization" in line):
        return ("", "", "Normalization", line, "")
    if len(parts) == 3:
        return (parts[0].strip(), parts[1].strip(), parts[2].strip(), "", "")
    return None


def write_rules_to_excel(rules_with_ref: list[tuple[str, str]], path: Path) -> None:
    """Write (rule_line, reference) pairs to an Excel file with columns Entity, Attribute, Normalization, Rule description, Few Shot examples, Reference."""
    import openpyxl
    from openpyxl.styles import Font

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Normalisation Rules"

    headers = ["Entity", "Attribute", "Normalization", "Rule description", "Few Shot examples", "Reference"]
    for col, h in enumerate(headers, start=1):
        cell = ws.cell(row=1, column=col, value=h)
        cell.font = Font(bold=True)

    row = 2
    for line, reference in rules_with_ref:
        parsed = _parse_rule_line(line)
        if parsed is None:
            continue
        entity, attr, norm, desc, few_shot = parsed
        ws.cell(row=row, column=1, value=entity)
        ws.cell(row=row, column=2, value=attr)
        ws.cell(row=row, column=3, value=norm)
        ws.cell(row=row, column=4, value=desc)
        ws.cell(row=row, column=5, value=few_shot)
        ws.cell(row=row, column=6, value=reference)
        row += 1

    # Column widths for readability
    ws.column_dimensions["A"].width = 14
    ws.column_dimensions["B"].width = 32
    ws.column_dimensions["C"].width = 16
    ws.column_dimensions["D"].width = 80
    ws.column_dimensions["E"].width = 60
    ws.column_dimensions["F"].width = 40

    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)


def write_rules_to_text(rules_with_ref: list[tuple[str, str]], path: Path) -> None:
    """Write (rule_line, reference) pairs to a plain text file with header Entity, Attribute, Normalization, Rule description, Few Shot examples, Reference."""
    path.parent.mkdir(parents=True, exist_ok=True)
    header = "Entity\tAttribute\tNormalization\tRule description\tFew Shot examples\tReference"
    out_lines = [header]
    for line, reference in rules_with_ref:
        parsed = _parse_rule_line(line)
        if parsed is None:
            continue
        entity, attr, norm, desc, few_shot = parsed
        out_lines.append(f"{entity}\t{attr}\t{norm}\t{desc}\t{few_shot}\t{reference}")
    path.write_text("\n".join(out_lines) + "\n", encoding="utf-8")
